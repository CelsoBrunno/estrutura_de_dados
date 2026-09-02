"""Visual da apostila — mesmo molde de Radar / Algoritmos."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

LARANJA = RGBColor(0xC4, 0x6B, 0x3A)
NAVY = RGBColor(0x0E, 0x24, 0x38)
TEXTO = RGBColor(0x1C, 0x1C, 0x1C)
CINZA = RGBColor(0x5C, 0x5C, 0x5C)
TEAL = RGBColor(0x1A, 0x6B, 0x64)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CAPA_SUB = RGBColor(0xC5, 0xD0, 0xD6)
COD_KW = RGBColor(0xE0, 0x8D, 0x5A)
COD_FN = RGBColor(0xE8, 0xC0, 0x7A)
COD_STR = RGBColor(0x8F, 0xBF, 0x7F)
COD_TXT = RGBColor(0xD6, 0xDE, 0xE4)
COD_CMT = RGBColor(0x7A, 0x8B, 0x99)

FILL_CAPA = "0E2438"
FILL_DICA = "F7EFE8"
FILL_ATENCAO = "F8EBE8"
FILL_OBJ = "E8F1F0"
FILL_PRATICA = "F3EEF8"
FILL_CODE = "1E2A32"

LARGURA_CEL = Twips(9638)
PASTA_IMAGENS = Path(__file__).resolve().parent / "imagens"

KEYWORDS = {
    "False", "None", "True", "and", "as", "assert", "break", "class",
    "continue", "def", "del", "elif", "else", "except", "finally", "for",
    "from", "global", "if", "import", "in", "is", "lambda", "nonlocal",
    "not", "or", "pass", "raise", "return", "try", "while", "with", "yield",
    "self",
}
BUILTINS = {
    "print", "len", "range", "list", "dict", "int", "str", "float", "bool",
    "super", "enumerate", "zip", "min", "max", "sum", "abs", "isinstance",
    "hasattr", "open", "input", "type", "sorted", "reversed", "set", "tuple",
}

_TOKEN = re.compile(
    r"(#.*)|"
    r"(f?(?:\"\"\"[\s\S]*?\"\"\"|'''[\s\S]*?'''|\"(?:\\.|[^\"\\])*\"|'(?:\\.|[^'\\])*'))|"
    r"(\b\w+\b)|"
    r"(\s+)|"
    r"(.)",
    re.MULTILINE,
)


def novo_documento() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(2)
    sec.header_distance = Cm(1.27)
    estilo_normal(doc)
    cabecalho(doc)
    rodape(doc)
    return doc


def estilo_normal(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)
    style.font.color.rgb = TEXTO
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.15


def slug(texto: str) -> str:
    nfkd = unicodedata.normalize("NFKD", texto)
    ascii_txt = "".join(c for c in nfkd if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", "_", ascii_txt.lower()).strip("_")
    if len(s) > 32:
        s = s[:32] + "_" + hashlib.md5(texto.encode()).hexdigest()[:6]
    return "s_" + s


def _rpr(run, nome, tamanho=None, bold=None, cor=None):
    run.font.name = nome
    r = run._element.get_or_add_rPr()
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.append(rFonts)
    rFonts.set(qn("w:ascii"), nome)
    rFonts.set(qn("w:hAnsi"), nome)
    rFonts.set(qn("w:eastAsia"), nome)
    if tamanho is not None:
        run.font.size = Pt(tamanho)
    if bold is not None:
        run.bold = bold
    if cor is not None:
        run.font.color.rgb = cor


def _preserve(run):
    t = run._r.find(qn("w:t"))
    if t is not None:
        t.set(qn("xml:space"), "preserve")


def esvaziar(p) -> None:
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def run_txt(p, texto, nome="Calibri", tamanho=12, bold=False, cor=TEXTO):
    r = p.add_run(texto)
    _rpr(r, nome, tamanho, bold, cor)
    if texto[:1] == " " or texto[-1:] == " " or "\t" in texto:
        _preserve(r)
    return r


_bookmark_id = 0


def bookmark(paragraph, nome: str) -> None:
    global _bookmark_id
    _bookmark_id += 1
    bid = str(_bookmark_id)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), nome)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _shade_tc(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "120"), ("left", "160"), ("bottom", "120"), ("right", "160")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _tabela_unica(doc: Document, fill: str):
    tabela = doc.add_table(rows=1, cols=1)
    tabela.autofit = True
    cell = tabela.cell(0, 0)
    cell.width = LARGURA_CEL
    _shade_tc(cell, fill)
    esvaziar(cell.paragraphs[0])
    return tabela, cell


def caixa(doc: Document, titulo: str, corpo: str, fill: str) -> None:
    _, cell = _tabela_unica(doc, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run_txt(p, titulo, "Trebuchet MS", 11, True, NAVY)
    if corpo:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        _texto_misto(p2, corpo)
    _espaco(doc, 8)


def dica(doc, texto): caixa(doc, "Dica", texto, FILL_DICA)
def atencao(doc, texto): caixa(doc, "Atenção", texto, FILL_ATENCAO)
def boa_pratica(doc, texto): caixa(doc, "Boa prática", texto, FILL_DICA)
def objetivo(doc, texto): caixa(doc, "Objetivo desta aula (2h30)", texto, FILL_OBJ)
def pratica(doc, titulo, texto): caixa(doc, titulo, texto, FILL_PRATICA)


def _texto_misto(p, texto: str, tamanho=12) -> None:
    partes = re.split(r"`([^`]+)`", texto)
    for i, parte in enumerate(partes):
        if not parte:
            continue
        if i % 2 == 1:
            run_txt(p, parte, "Consolas", tamanho, False, TEAL)
        else:
            run_txt(p, parte, "Calibri", tamanho, False, TEXTO)


def corpo(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    _texto_misto(p, texto)


def bullets(doc: Document, itens: list[str]) -> None:
    for item in itens:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.4)
        run_txt(p, "▸  ", "Calibri", 12, False, LARANJA)
        _texto_misto(p, item)


def h_rotulo(doc: Document, texto: str, bookmark_nome: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    run_txt(p, texto, "Trebuchet MS", 9, True, LARANJA)
    if bookmark_nome:
        bookmark(p, bookmark_nome)


def h_titulo(doc: Document, texto: str, bookmark_nome: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    run_txt(p, texto, "Trebuchet MS", 22, True, NAVY)
    if bookmark_nome:
        bookmark(p, bookmark_nome)


def h_aula(doc: Document, numero: int, titulo: str) -> None:
    quebra(doc)
    h_rotulo(doc, f"AULA {numero}", slug(f"AULA {numero}"))
    h_titulo(doc, titulo, slug(titulo))


def h_secao(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run_txt(p, texto, "Trebuchet MS", 15, True, TEAL)
    bookmark(p, slug(texto))


def h_antes(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run_txt(p, "Antes da próxima aula, você precisa conseguir", "Trebuchet MS", 12, True, NAVY)


def figura(doc: Document, arquivo: str, legenda: str, largura_cm: float = 15.2) -> None:
    caminho = PASTA_IMAGENS / arquivo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(caminho), width=Cm(largura_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    run_txt(cap, legenda, "Calibri", 9, False, CINZA)


def caption(doc: Document, nome_arquivo: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run_txt(p, nome_arquivo, "Trebuchet MS", 9, True, CINZA)


def _cor_token(tipo: str, texto: str, depois_de_def: bool) -> RGBColor:
    if tipo == "cmt":
        return COD_CMT
    if tipo == "str":
        return COD_STR
    if tipo == "id":
        if texto in KEYWORDS:
            return COD_KW
        if depois_de_def or texto in BUILTINS:
            return COD_FN
        return COD_TXT
    return COD_TXT


def destacar(p, codigo: str) -> None:
    depois_de_def = False
    for m in _TOKEN.finditer(codigo):
        cmt, string, ident, espaco, outro = m.groups()
        if cmt:
            run_txt(p, cmt, "Consolas", 10, False, COD_CMT)
            depois_de_def = False
        elif string:
            run_txt(p, string, "Consolas", 10, False, COD_STR)
            depois_de_def = False
        elif ident:
            cor = _cor_token("id", ident, depois_de_def)
            run_txt(p, ident, "Consolas", 10, False, cor)
            depois_de_def = ident in ("def", "class")
        elif espaco:
            run_txt(p, espaco, "Consolas", 10, False, COD_TXT)
        else:
            run_txt(p, outro or "", "Consolas", 10, False, COD_TXT)
            if outro not in (" ",):
                depois_de_def = False


def codigo(doc: Document, fonte: str, arquivo: str | None = None) -> None:
    if arquivo:
        caption(doc, arquivo)
    _, cell = _tabela_unica(doc, FILL_CODE)
    linhas = fonte.replace("\t", "    ").splitlines() or [""]
    first = True
    for linha in linhas:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        destacar(p, linha if linha else " ")
    _espaco(doc, 8)


def _espaco(doc: Document, pts: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    p.paragraph_format.space_before = Pt(0)


def quebra(doc: Document) -> None:
    doc.add_page_break()


def _campo(p, instrucao: str, preview: str, cor=None):
    r1 = p.add_run()
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc1)

    r2 = p.add_run()
    ins = OxmlElement("w:instrText")
    ins.set(qn("xml:space"), "preserve")
    ins.text = f" {instrucao} "
    r2._r.append(ins)

    r3 = p.add_run()
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc2)

    r4 = p.add_run(preview)
    if cor:
        _rpr(r4, "Calibri", 8, False, cor)
    else:
        _rpr(r4, "Calibri", 12, False, CINZA)

    r5 = p.add_run()
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "end")
    r5._r.append(fc3)


def cabecalho(doc: Document) -> None:
    p = doc.sections[0].header.paragraphs[0]
    esvaziar(p)
    run_txt(p, "PYTHON", "Trebuchet MS", 8, True, LARANJA)
    run_txt(p, "    Apostila passo a passo  ·  Estruturas de Dados", "Calibri", 8, False, CINZA)


def rodape(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    esvaziar(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_txt(
        p,
        "Celso Brunno Rocha Custódio de Campos  ·  Estruturas de Dados  ·  ",
        "Calibri",
        8,
        False,
        CINZA,
    )
    _campo(p, "PAGE", "1", LARANJA)


def capa(doc: Document) -> None:
    _, cell = _tabela_unica(doc, FILL_CAPA)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(6)
    run_txt(p, "CURSO DE PYTHON", "Trebuchet MS", 11, True, LARANJA)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    run_txt(p2, "Estruturas de Dados", "Trebuchet MS", 32, True, BRANCO)

    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(10)
    run_txt(
        p3,
        "Apostila passo a passo: da memória RAM ao projeto final",
        "Calibri",
        13,
        False,
        CAPA_SUB,
    )

    p4 = cell.add_paragraph()
    p4.paragraph_format.space_after = Pt(16)
    run_txt(p4, "Python 3   ·   VS Code   ·   estruturas do zero", "Trebuchet MS", 10, False, LARANJA)

    p5 = cell.add_paragraph()
    p5.paragraph_format.space_after = Pt(2)
    run_txt(p5, "APOSTILA ELABORADA POR", "Trebuchet MS", 8, True, LARANJA)

    p6 = cell.add_paragraph()
    p6.paragraph_format.space_after = Pt(8)
    run_txt(
        p6,
        "Professor Celso Brunno Rocha Custódio de Campos",
        "Trebuchet MS",
        14,
        True,
        BRANCO,
    )
    _espaco(doc, 10)


def linha_sumario(doc: Document, rotulo: str | None, titulo: str, ancora: str, nivel: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(4 if nivel == 0 else 0)
    p.paragraph_format.left_indent = Cm(0.6 * nivel)
    tab = p.paragraph_format.tab_stops.add_tab_stop(
        Cm(17), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    _ = tab
    if rotulo:
        run_txt(p, rotulo + "  ", "Trebuchet MS", 8, True, LARANJA)
    run_txt(p, titulo, "Calibri", 12 if nivel == 0 else 11, nivel == 0, TEXTO if nivel == 0 else CINZA)
    run_txt(p, "\t", "Calibri", 12, False, CINZA)
    _campo(p, f"PAGEREF {ancora} \\h", "·", CINZA)
